"""Fill the SEC cover-letter form for the Wikes & Hossain detection manuscript.

Personal contact fields (address, phone, co-author email, signatures, dates) are
left as highlighted [ADD ...] markers for the owner; they are not stored here.

Usage: python3 scripts/fill_sec_cover_letter.py <blank_form.docx> <out.docx>
"""
import sys

import docx
from docx.enum.text import WD_COLOR_INDEX

TITLE = ("Socioeconomic Challenges of AI-Drafted Decision Records: "
         "Can Expert Reviewers Detect Unsupported Conclusions?")
AUTHORS = [
    {"name": "Phillip Wikes", "position": "Independent Researcher",
     "email": "info@jrsstandard.com", "orcid": "0009-0009-7230-3661",
     "roles": {"Conceptualisation", "Methodology", "Data curation", "Formal analysis", "Investigation",
               "Project administration", "Resources", "Writing – original draft", "Writing – review & editing"}},
    {"name": "Ubayet Hossain, FRM", "position": "Independent Financial Risk & Model Validation Professional",
     "email": "[ADD: his own email address]", "orcid": "0009-0001-5304-2170",
     "roles": {"Methodology", "Validation", "Formal analysis", "Supervision", "Writing – review & editing"}},
]


def put(par, text):
    r = par.add_run(" " + text)
    if text.startswith("["):
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW


def main(src, out):
    d = docx.Document(src)
    body = d.element.body
    els = list(body.iterchildren())
    paras = {id(p._p): p for p in d.paragraphs}
    tables = {id(t._tbl): t for t in d.tables}

    def P(i):
        return paras[id(els[i])]

    def T(i):
        return tables[id(els[i])]

    for i in (3, 108):
        p = P(i)
        for r in p.runs:
            for blank in ("_________________", "_______"):
                if blank in r.text:
                    r.text = r.text.replace(blank, TITLE)
    for start, tbl, a in ((8, 19, AUTHORS[0]), (21, 32, AUTHORS[1])):
        put(P(start + 1), a["position"] + "; no institutional affiliation")
        put(P(start + 2), "[ADD BEFORE SENDING]")
        put(P(start + 3), "Not applicable")
        put(P(start + 4), a["email"])
        put(P(start + 5), "[ADD BEFORE SENDING]")
        put(P(start + 6), a["orcid"])
        put(P(start + 7), "[ADD IF AVAILABLE]")
        put(P(start + 8), "[DATE]")
        P(start).add_run(" " + a["name"]).bold = True
        for row in T(tbl).rows:
            cells = row.cells
            for c in range(0, len(cells) - 1, 2):
                role = cells[c].text.strip().replace("-", "–")
                if role in a["roles"]:
                    cells[c + 1].text = "X"
    sig = T(128)
    sig.rows[1].cells[0].text = AUTHORS[0]["name"]
    sig.rows[2].cells[0].text = AUTHORS[1]["name"]
    for row in list(sig.rows)[3:]:
        row._tr.getparent().remove(row._tr)
    T(133).rows[1].cells[0].text = AUTHORS[0]["name"]
    # Remove unused Author 3-5 and Contributor blocks (elements 33..105).
    for el in els[33:106]:
        body.remove(el)
    d.save(out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
