# Builds the blind basis-note re-coding sheet for the SLGR submission (2026-10-07), and its answer key.
# Usage: python3 scripts/build_note_recode_sheet.py <out_dir> research/JCI_SUBMISSION_2026-08-28/02_DATA/JCI_JRS_Construct_Coding_Frame.csv
# The shuffle uses a fixed seed, so the key can be regenerated exactly. Send only the sheet to the coder.
import csv, random, sys, hashlib
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
OUT, FRAME = sys.argv[1], sys.argv[2]
rows = list(csv.DictReader(open(FRAME)))
assert len(rows) == 24
order = rows[:]; random.Random(20261007).shuffle(order)       # fixed seed: the order is reproducible
key = [(f'N{i+1:02d}', r['Case ID']) for i, r in enumerate(order)]
d = Document()
st = d.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(11); st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
for m in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'): setattr(d.sections[0], m, Inches(0.8))
d.add_heading('Basis-note coding sheet', level=1)
for s in [
 'This sheet asks one question about each of 24 basis notes. They are the notes you wrote when you made each read. The notes are shown in a shuffled order, with no case number, no read and no earlier coding, so that each answer depends only on the words of the note.',
 'The rule. Does the note explicitly state that the underlying basis could not be rebuilt, reconstructed or recreated from the source? Answer Yes only if the note says so in words. Answer No if it does not, even if you think it is implied. Inference does not count.',
 'How to answer. Put Yes or No in the Answer column for every note. If a note is genuinely unclear, answer anyway and add a few words in the Comment column. Please work alone, do not look at the case records, the manuscript or any earlier coding while you do this, and do not change an answer after you finish the sheet.',
 'When you are done, add your name and the date at the bottom and return the sheet. It takes about 30 to 45 minutes.',
]: d.add_paragraph(s)
t = d.add_table(rows=1, cols=4); t.style = 'Table Grid'
for c, h, w in zip(t.rows[0].cells, ['No.', 'Basis note (verbatim)', 'Answer (Yes / No)', 'Comment'], [0.5, 4.3, 0.9, 1.2]):
    c.text = h; c.paragraphs[0].runs[0].bold = True; c.width = Inches(w)
for (n, _), r in zip(key, order):
    cells = t.add_row().cells
    for c, v, w in zip(cells, [n, r['Supporting Note'].strip(), '', ''], [0.5, 4.3, 0.9, 1.2]): c.text = v; c.width = Inches(w)
d.add_paragraph('')
d.add_paragraph('Coder name: ______________________        Date completed: ______________')
d.add_paragraph('I confirm I coded these notes alone, without looking at the reads, the case records, the manuscript or any earlier coding.  Initials: ______')
d.save(OUT + '/Basis_Note_Coding_Sheet.docx')
with open(OUT + '/Basis_Note_Coding_KEY_do_not_send.csv', 'w', newline='') as f:
    w = csv.writer(f); w.writerow(['Sheet No.', 'Case ID', 'Read', 'Earlier code (AI-applied, 2026-08-08)'])
    by = {r['Case ID']: r for r in rows}
    for n, cid in key: w.writerow([n, cid, by[cid]['JRS Read'], by[cid]['Reconstructability Failure Explicitly Stated']])
print('ok', hashlib.sha256(open(OUT + '/Basis_Note_Coding_Sheet.docx', 'rb').read()).hexdigest()[:12])
