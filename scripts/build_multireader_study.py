#!/usr/bin/env python3
"""Multi-reader extension of the 32-case public-records study (SLGR submission, 2026-10-07).

Design: four independent volunteer readers (A to D). The 32 cases are split into two halves,
stratified by the original read so each half carries the same mix. Readers A and C read half 1,
B and D read half 2, so every case receives two new blind reads in addition to the original.
Each reader sees a different shuffled order. Every assignment uses a fixed seed and is reproducible.

Usage:
  python3 scripts/build_multireader_study.py packets <out_dir>   write reader packets and the private key
  python3 scripts/build_multireader_study.py score <returns.csv>  score returned reads (see score_reads)

The read question and the three definitions are the exact wording the first second reader used
(recheck.html), so the new reads are comparable with the existing ten.
"""
import csv, random, sys, os

DATA = 'research/JCI_SUBMISSION_2026-08-28/02_DATA/JCI_JRS_32_Case_Master_Dataset.csv'
QUESTION = ('Could a reviewer with no access to the author, no institutional memory, and no supplementary '
            'explanation identify the basis for each conclusion from the record alone?')
DEFS = [('Ready', 'The basis for each conclusion is present in the record and can be rebuilt from it without asking anyone.'),
        ('Needs work', 'The record states conclusions whose basis is partly present, so a cold reader could reconstruct some of it and would have to assume the rest.'),
        ('Gap', 'The record does not carry a reconstructable basis at all. A cold reader could not say why the conclusion was reached.')]
SEED = 20261007

def halves(rows):
    rng = random.Random(SEED); h1, h2 = [], []
    for read in ('Ready', 'Needs work', 'Gap'):
        grp = sorted([r for r in rows if r['JRS Read'] == read], key=lambda r: r['Case ID']); rng.shuffle(grp)
        for i, r in enumerate(grp): (h1 if i % 2 == 0 else h2).append(r)
    return h1, h2

def assignments(rows):
    h1, h2 = halves(rows); out = {}
    for reader, half in (('A', h1), ('B', h2), ('C', h1), ('D', h2)):
        order = half[:]; random.Random(SEED + ord(reader)).shuffle(order); out[reader] = order
    return out

def packets(out_dir):
    from docx import Document
    from docx.shared import Pt, Inches
    rows = list(csv.DictReader(open(DATA))); assert len(rows) == 32
    os.makedirs(out_dir, exist_ok=True)
    key = []
    for reader, order in assignments(rows).items():
        d = Document(); st = d.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(11)
        for m in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'): setattr(d.sections[0], m, Inches(0.8))
        d.add_heading('Independent read: %d public-records cases (Reader %s)' % (len(order), reader), level=1)
        for s in ['Thank you for helping. You will read %d public documents and answer one question about each. You are not checking anyone\'s work and you will not see anyone else\'s answers. There is no right total.' % len(order),
                  'The question: ' + QUESTION]:
            d.add_paragraph(s)
        for name, text in DEFS:
            p = d.add_paragraph(); p.add_run(name + '. ').bold = True; p.add_run(text)
        d.add_paragraph('Ground rules. Open the link and read the source itself before deciding. Do not look up how a case came out; if you already know, tick the box and answer anyway. '
                        'Do not discuss the cases with anyone else in the study until you have returned the sheet. If a link does not open, leave that case blank and say so. '
                        'Give a short reason for each answer, one or two sentences in your own words. Expect about three hours; you can stop and come back.')
        t = d.add_table(rows=1, cols=5); t.style = 'Table Grid'
        for c, h in zip(t.rows[0].cells, ['No.', 'Source and link', 'Read (Ready / Needs work / Gap)', 'Reason', 'Knew outcome? (Y/N)']):
            c.text = h; c.paragraphs[0].runs[0].bold = True
        for i, r in enumerate(order, 1):
            no = '%s%02d' % (reader, i); key.append([no, reader, r['Case ID']])
            cells = t.add_row().cells
            cells[0].text = no
            cells[1].text = '%s, %s. %s\n%s' % (r['Source type'], r['Decision/source year'], r['Citation'], r['Public URL'])
        d.add_paragraph('')
        d.add_paragraph('Name: ______________________   Professional background (one line): ______________________   Date returned: __________')
        d.add_paragraph('Prior familiarity with this study or its authors (one line, or "none"): ______________________')
        d.add_paragraph('I read these cases alone and did not see the original reads or anyone else\'s answers.  Initials: ______')
        d.save(os.path.join(out_dir, 'Reader_%s_Packet.docx' % reader))
    with open(os.path.join(out_dir, 'MULTIREADER_KEY_do_not_send.csv'), 'w', newline='') as f:
        csv.writer(f).writerows([['Packet No.', 'Reader', 'Case ID']] + key)
    print('wrote 4 packets and the key to', out_dir)

# ---- scoring ----------------------------------------------------------------------------------
ORD = {'Ready': 0, 'Needs work': 1, 'Gap': 2}

def kalpha(units, metric):
    """Krippendorff's alpha. units: list of lists of values (one list per case, any number of readers)."""
    vals = sorted({v for u in units for v in u})
    o = {(a, b): 0.0 for a in vals for b in vals}
    for u in units:
        m = len(u)
        if m < 2: continue
        for i in range(m):
            for j in range(m):
                if i != j: o[(u[i], u[j])] += 1.0 / (m - 1)
    n_c = {c: sum(o[(c, k)] for k in vals) for c in vals}; n = sum(n_c.values())
    if metric == 'nominal':
        d = lambda a, b: 0.0 if a == b else 1.0
    else:  # ordinal
        def d(a, b):
            lo, hi = sorted((a, b)); s = sum(n_c[g] for g in vals if lo <= g <= hi) - (n_c[lo] + n_c[hi]) / 2.0
            return s * s
    Do = sum(o[(a, b)] * d(a, b) for a in vals for b in vals) / n
    De = sum(n_c[a] * n_c[b] * d(a, b) for a in vals for b in vals) / (n * (n - 1))
    return 1.0 - Do / De if De else float('nan')

def score_reads(path):
    """returns.csv columns: Case ID, Reader, Read. The original read is added as reader 'P'."""
    base = {r['Case ID']: r['JRS Read'] for r in csv.DictReader(open(DATA))}
    by = {}
    for r in csv.DictReader(open(path)):
        if r['Read'].strip() in ORD: by.setdefault(r['Case ID'], []).append(ORD[r['Read'].strip()])
    units_new = [v for v in by.values()]
    units_all = [[ORD[base[c]]] + by.get(c, []) for c in base]
    rng = random.Random(SEED)
    def boot(units, metric, B=4000):
        est = []
        for _ in range(B):
            s = [units[rng.randrange(len(units))] for _ in units]
            a = kalpha(s, metric)
            if a == a: est.append(a)
        est.sort(); return est[int(0.025 * len(est))], est[int(0.975 * len(est)) - 1]
    for label, U in (('new readers only', units_new), ('new readers plus the original read', units_all)):
        for metric in ('ordinal', 'nominal'):
            a = kalpha(U, metric); lo, hi = boot(U, metric)
            print('%-38s %-8s alpha = %.3f  (95%% bootstrap %.3f to %.3f)' % (label, metric, a, lo, hi))
    pairs = [(ORD[base[c]], v) for c, vs in by.items() for v in vs]
    print('agreement of each new read with the original read: %d of %d' % (sum(a == b for a, b in pairs), len(pairs)))

if __name__ == '__main__':
    if sys.argv[1] == 'packets': packets(sys.argv[2])
    elif sys.argv[1] == 'score': score_reads(sys.argv[2])
